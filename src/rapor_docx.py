"""docs/rapor_metni_stereo.md dosyasini Word belgesine cevir.

Resmi sablonun bicim kurallarini uygular:
  bolum basligi    14 punto, kalin, BUYUK HARF
  alt bolum        12 punto, kalin
  govde            11 punto, 1.15 satir araligi
  sekil/cizelge
  altyazisi        11 punto, kalin, ORTALI
  sekiller         satir ortali, metne gomulu degil

[SEKIL N BURAYA] isaretlerinin yerine gercek resmi koyar ve altina
altyazisini yazar. Markdown tablolarini Word tablosuna cevirir.

Cikti bir TASLAKTIR - kapak, kurum onayi, degerlendirme formu,
icindekiler ve sayfa altligi Word'de elle eklenecek (sablon dosyasi
zaten bunlari iceriyor).

Kullanim:
    python src/rapor_docx.py
"""
from __future__ import annotations

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, RGBColor

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GIRDI = os.path.join(KOK, "docs", "rapor_metni_stereo.md")
CIKTI = os.path.join(KOK, "output", "reports", "rapor_stereo_bolumu.docx")

GOVDE_PT = 11
BOLUM_PT = 14
ALT_PT = 12
SATIR_ARALIGI = 1.15


def govde_bicimle(belge):
    st = belge.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(GOVDE_PT)
    st.paragraph_format.line_spacing = SATIR_ARALIGI
    st.paragraph_format.space_after = Pt(6)


def kalin_isaretleri_uygula(p, metin):
    """**kalin** isaretlerini gercek kalin yaziya cevir."""
    for i, parca in enumerate(re.split(r"\*\*(.+?)\*\*", metin)):
        if not parca:
            continue
        r = p.add_run(parca.replace("`", ""))
        r.bold = (i % 2 == 1)


def altyazi_ekle(belge, metin):
    p = belge.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run(metin)
    r.bold = True
    r.font.size = Pt(GOVDE_PT)
    r.font.color.rgb = RGBColor(0, 0, 0)


def tablo_ekle(belge, satirlar):
    """Markdown tablosunu Word tablosuna cevir."""
    hucreler = [[h.strip() for h in s.strip().strip("|").split("|")]
                for s in satirlar]
    # ikinci satir ayirac (---|---) - at
    if len(hucreler) > 1 and all(set(h) <= set("-: ") for h in hucreler[1]):
        basliklar, govde = hucreler[0], hucreler[2:]
    else:
        basliklar, govde = hucreler[0], hucreler[1:]
    t = belge.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, b in enumerate(basliklar):
        hc = t.rows[0].cells[i]
        hc.text = ""
        p = hc.paragraphs[0]
        kalin_isaretleri_uygula(p, b)
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(GOVDE_PT - 1)
    for satir in govde:
        hs = t.add_row().cells
        for i, d in enumerate(satir[:len(basliklar)]):
            hs[i].text = ""
            p = hs[i].paragraphs[0]
            kalin_isaretleri_uygula(p, d)
            for r in p.runs:
                r.font.size = Pt(GOVDE_PT - 1)
    return t


def main():
    if not os.path.exists(GIRDI):
        print(f"Girdi yok: {GIRDI}")
        return 1
    metin = open(GIRDI, encoding="utf-8").read()
    # yalnizca rapor govdesini al - basindaki aciklama ve sonundaki
    # kontrol listesi Word'e girmemeli
    bas = metin.index("# 4. STEREO KAMERA")
    son = metin.index("## SEKIL VE CIZELGE YERLESIM OZETI")
    govde = metin[bas:son]

    belge = Document()
    for b in belge.sections:
        b.left_margin = b.right_margin = Cm(2.5)
    govde_bicimle(belge)

    satirlar = govde.split("\n")
    i = 0
    sekil_no = 0
    eklenen_sekil, eklenen_cizelge = 0, 0
    while i < len(satirlar):
        s = satirlar[i]
        d = s.strip()

        # --- sekil yer tutucusu
        m = re.match(r"\*\*\[SEKIL (\d+) BURAYA\]\*\*", d)
        if m:
            sekil_no = int(m.group(1))
            dosya = alt = None
            j = i + 1
            while j < len(satirlar) and satirlar[j].strip():
                if satirlar[j].startswith("Dosya:"):
                    dosya = satirlar[j].split("`")[1]
                if satirlar[j].startswith("Altyazi:"):
                    alt = satirlar[j].split("Altyazi:")[1].strip()
                j += 1
            yol = os.path.join(KOK, dosya) if dosya else None
            if yol and os.path.exists(yol):
                p = belge.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(yol, width=Cm(15))
                eklenen_sekil += 1
            else:
                p = belge.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run(f"[Sekil {sekil_no} — dosya bulunamadi: "
                          f"{dosya}]").italic = True
            if alt:
                altyazi_ekle(belge, alt.replace("**", ""))
            i = j
            continue

        # --- cizelge yer tutucusu: sonraki tabloyu bekle
        if re.match(r"\*\*\[CIZELGE \d+ BURAYA\]\*\*", d):
            j = i + 1
            while j < len(satirlar) and not satirlar[j].strip().startswith("|"):
                j += 1
            tablo = []
            while j < len(satirlar) and satirlar[j].strip().startswith("|"):
                tablo.append(satirlar[j])
                j += 1
            if tablo:
                tablo_ekle(belge, tablo)
                eklenen_cizelge += 1
            while j < len(satirlar) and not satirlar[j].strip().startswith(
                    "Altyazi:"):
                if satirlar[j].strip() and not satirlar[j].startswith("|"):
                    break
                j += 1
            if j < len(satirlar) and satirlar[j].strip().startswith("Altyazi:"):
                altyazi_ekle(
                    belge,
                    satirlar[j].split("Altyazi:")[1].strip().replace("**", ""))
                j += 1
            i = j
            continue

        # --- basliklar
        if d.startswith("# "):
            p = belge.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            r = p.add_run(d[2:].upper())
            r.bold = True
            r.font.size = Pt(BOLUM_PT)
            i += 1
            continue
        if d.startswith("## "):
            p = belge.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            r = p.add_run(d[3:])
            r.bold = True
            r.font.size = Pt(ALT_PT)
            i += 1
            continue

        # --- kod blogu (formuller)
        if d.startswith("    ") or s.startswith("    "):
            p = belge.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(d)
            r.font.name = "Consolas"
            r.font.size = Pt(GOVDE_PT)
            i += 1
            continue

        # --- madde
        if d.startswith("- "):
            p = belge.add_paragraph(style="List Bullet")
            kalin_isaretleri_uygula(p, d[2:])
            i += 1
            continue

        if not d or d == "---":
            i += 1
            continue

        # --- normal paragraf: bos satira kadar birlestir
        parca = []
        while i < len(satirlar) and satirlar[i].strip() \
                and not satirlar[i].strip().startswith(("#", "|", "- ",
                                                        "**[", "Dosya:",
                                                        "Altyazi:")):
            parca.append(satirlar[i].strip())
            i += 1
        if parca:
            p = belge.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            kalin_isaretleri_uygula(p, " ".join(parca))
        else:
            i += 1

    os.makedirs(os.path.dirname(CIKTI), exist_ok=True)
    belge.save(CIKTI)
    print(f"Olusturuldu: {os.path.relpath(CIKTI, KOK)}")
    print(f"  {eklenen_sekil} sekil, {eklenen_cizelge} cizelge eklendi")
    print()
    print("Word'de elle yapilacaklar:")
    print("  - Kapak, kurum onayi ve degerlendirme formu sayfalari")
    print("  - Sayfa altligina ogrenci ad-soyad ve numara")
    print("  - Icindekiler / Sekiller / Cizelgeler listeleri")
    print("  - Birinci projeye gore bolum ve sekil numaralarini kaydir")
    return 0


if __name__ == "__main__":
    sys.exit(main())
